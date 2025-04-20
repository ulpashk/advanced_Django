from celery import shared_task
import spacy
from pymongo import MongoClient
from bson import ObjectId
from pdfminer.high_level import extract_text as extract_pdf_text
from docx import Document
from django.apps import apps
import os
from django.conf import settings

nlp = spacy.load("en_core_web_sm")

mongo_client = MongoClient(settings.MONGO_URI)
mongo_db = mongo_client[settings.MONGO_DB_NAME]
parsed_collection = mongo_db["resumes"]

def extract_text(file_path):
    if file_path.endswith(".pdf"):
        return extract_pdf_text(file_path)
    elif file_path.endswith(".docx"):
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs])
    return ""

def parse_resume(text):
    doc = nlp(text)
    skills = []
    education = []
    experience = []

    for ent in doc.ents:
        if ent.label_ in ["ORG", "GPE"]:
            experience.append(ent.text)
        elif ent.label_ == "EDUCATION":
            education.append(ent.text)

    # Sample skill list
    skill_keywords = {"python", "django", "sql", "javascript", "react", "aws"}
    for token in doc:
        if token.text.lower() in skill_keywords:
            skills.append(token.text)

    return {
        "skills": list(set(skills)),
        "education": list(set(education)),
        "experience": list(set(experience)),
    }

@shared_task
def parse_resume_task(resume_id):
    try:
        ResumeUpload = apps.get_model('resumes', 'ResumeUpload')  # Dynamically get the model
        resume = ResumeUpload.objects.get(id=resume_id)
        resume.status = "PROCESSING"
        resume.save()

        file_path = resume.file.path
        text = extract_text(file_path)
        parsed_data = parse_resume(text)

        result = parsed_collection.insert_one(parsed_data)
        resume.parsed_data_id = str(result.inserted_id)
        resume.status = "COMPLETED"
        resume.save()
    except Exception as e:
        resume.status = "FAILED"
        resume.save()
        print(f"Resume parsing failed: {e}")




























# from celery import shared_task
# import spacy
# from pymongo import MongoClient
# from bson import ObjectId
# from pdfminer.high_level import extract_text as extract_pdf_text
# from docx import Document
# from resumes.models import ResumeUpload
# import os
# from django.conf import settings


# nlp = spacy.load("en_core_web_sm")


# mongo_client = MongoClient(settings.MONGO_URI)
# mongo_db = mongo_client[settings.MONGO_DB_NAME]
# parsed_collection = mongo_db["resumes"]


# def extract_text(file_path):
#     if file_path.endswith(".pdf"):
#         return extract_pdf_text(file_path)
#     elif file_path.endswith(".docx"):
#         doc = Document(file_path)
#         return "\n".join([p.text for p in doc.paragraphs])
#     return ""

# def parse_resume(text):
#     doc = nlp(text)
#     skills = []
#     education = []
#     experience = []

#     for ent in doc.ents:
#         if ent.label_ in ["ORG", "GPE"]:
#             experience.append(ent.text)
#         elif ent.label_ == "EDUCATION":
#             education.append(ent.text)

#     # Sample skill list
#     skill_keywords = {"python", "django", "sql", "javascript", "react", "aws"}
#     for token in doc:
#         if token.text.lower() in skill_keywords:
#             skills.append(token.text)

#     return {
#         "skills": list(set(skills)),
#         "education": list(set(education)),
#         "experience": list(set(experience)),
#     }

# @shared_task
# def parse_resume_task(resume_id):
#     try:
#         resume = ResumeUpload.objects.get(id=resume_id)
#         resume.status = "PROCESSING"
#         resume.save()

#         file_path = resume.file.path
#         text = extract_text(file_path)
#         parsed_data = parse_resume(text)

#         result = parsed_collection.insert_one(parsed_data)
#         resume.parsed_data_id = str(result.inserted_id)
#         resume.status = "COMPLETED"
#         resume.save()
#     except Exception as e:
#         resume.status = "FAILED"
#         resume.save()
#         print(f"Resume parsing failed: {e}")






# from celery import shared_task
# from .models import ResumeUpload
# from django.conf import settings
# from pymongo import MongoClient
# import os
# import uuid


# @shared_task
# def parse_resume_task(resume_id):
#     resume = None
#     try:
#         resume = ResumeUpload.objects.get(id=resume_id)
#         resume.status = "PROCESSING"
#         resume.save()

#         # Parse logic here - dummy example
#         parsed = {
#             "name": "Han Sohi",
#             "email": "sohi@example.com",
#             "skills": ["Python", "React"],
#             "education": "BSc Computer Science",
#             "experience": "3 years at XYZ Corp"
#         }

#         # Store in MongoDB
#         client = MongoClient(settings.MONGO_URI)
#         db = client[settings.MONGO_DB_NAME]
#         collection = db["resumes"]
#         result = collection.insert_one(parsed)
#         # result = db.resumes.insert_one(parsed)

#         resume.status = "COMPLETED"
#         resume.parsed_data_id = str(result.inserted_id)
#         resume.save()

#     except Exception as e:
#         if resume:
#             resume.status = "FAILED"
#             resume.save()
#         print(f"Resume parsing failed: {e}")







# @shared_task
# def parse_resume_task(resume_id):
#     try:
#         resume = ResumeUpload.objects.get(id=resume_id)
#         resume.status = "PROCESSING"
#         resume.save()

#         # Parse logic here - dummy example
#         parsed = {
#             "name": "Han Sohi",
#             "email": "sohi@example.com",
#             "skills": ["Python", "React"],
#             "education": "BSc Computer Science",
#             "experience": "3 years at XYZ Corp"
#         }

#         # Store in MongoDB
#         client = MongoClient(settings.MONGO_URI)
#         db = client[settings.MONGO_DB_NAME]
#         result = db.resumes.insert_one(parsed)
        
#         resume.status = "COMPLETED"
#         resume.parsed_data_id = str(result.inserted_id)
#         resume.save()

#     except Exception as e:
#         resume.status = "FAILED"
#         resume.save()
#         print(f"Resume parsing failed: {e}")


